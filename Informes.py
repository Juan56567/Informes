from docx import Document
import os
from Enunciados import *


location = "C:/Users/Lenovo/Downloads/CILINDROS"

datos = os.listdir(location)


print("Cargando")

counter = 0
for i in datos:
    temp = location + "/" + i
    doc = Document(temp)
    tables = doc.tables

    camisa = tables[1].cell(26, 11)
    vastago = tables[1].cell(26, 25)
    piston = tables[2].cell(3, 0)
    montajeF = tables[2].cell(3, 1)
    montajeP = tables[2].cell(7, 0)
    tapa = tables[2].cell(7, 1)
    fecha = tables[1].cell(4,25)

    #CAMBIAR EL FORMATO DE LA FECHA
    ftemp = fecha.text.split(",")
    fecha.text = ftemp[0]

    for i in range(10, 18):
        read = tables[1].cell(i, 25)
        if len(read.text) < 2:
            read.text = "BUENO"


    #AGREGAR EN UN LISTA TODAS LAS CELDAS A CAMBIAR
    componentes = [camisa, vastago, piston, montajeF, montajeP, tapa]

    for c in componentes:
        text = c.text.split("-")
        if len(text) > 1:
            if text[0] == "BUENO":
                c.text = bueno
            elif text[0] == "CAMBIO":
                if "CORROSION" in text[1]:
                    c.text = cambioCorrosion
                elif "TOLERANCIA" in text[1]:
                    c.text = cambioTolerancia
                elif "GOLPES" in text[1]:
                    c.text = cambioGolpes
                elif "DESPRENDIMIENTO" in text[1]:
                    c.text = cambioDesprendimiento
                elif "DESGASTE" in text[1]:
                    c.text = cambioDesgaste
                elif "RALLAS" in text[1]:
                    c.text = cambioRallas
                elif "FLEXION" in text[1]:
                    c.text = cambioFlexion
                else:
                    pass
            else:
                if "CROMAR" in text[1]:
                    c.text = rectiCromo
                elif "BRILLAR" in text[1]:
                    c.text = rectiBrillar
                elif "BRUÑIR" in text[1]:
                    c.text = rectiBrunir
                elif "GOLPES" in text[1]:
                    c.text = rectiDeformacion
                elif "ROSCA" in text[1]:
                    c.text = rectiRosca
        else:
            pass
    doc.save(temp)

    counter += 1
    print(counter)

print("Informes terminados")
